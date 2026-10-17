"""Prepare a client workbook and product photographs without changing originals.

Run with the bundled Python runtime (openpyxl and Pillow are required).
The manifest is the only input used by the controlled Supabase import step.
"""

import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
import re
import shutil
import unicodedata
import zipfile

import openpyxl
from PIL import Image, ImageDraw, ImageOps


def key(value):
    return re.sub(r"[^A-Z0-9]", "", unicodedata.normalize("NFKD", str(value)).upper())


def slug(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def sql_literal(value):
    return "'" + value.replace("'", "''") + "'"


def prepare_sql(output, manifest):
    """Generate one atomic import, but do not connect or execute anything."""
    rows = [r for r in manifest["ready"] if r["match"] == "existing"]
    images = [r for r in manifest["images"] if r["model"] in {p["model"] for p in rows}]
    payload = sql_literal(json.dumps(rows, ensure_ascii=False))
    photos = sql_literal(json.dumps(images, ensure_ascii=False))
    source_table = f"""create temporary table import_source on commit drop as
select * from jsonb_to_recordset({payload}::jsonb)
as x("sourceRow" integer, name text, brand text, "sourceModel" text, color text,
     specifications text, model text, sku text, "existingSourceRef" text,
     "existingSlug" text, category text, match text);

create temporary table import_targets on commit drop as
select s.*, v.id as variant_id, v.product_id
from import_source s
join public.catalog_variants v on
  regexp_replace(upper(v.manufacturer_model), '[^A-Z0-9]', '', 'g') =
  regexp_replace(upper(s.model), '[^A-Z0-9]', '', 'g')
join public.catalog_products p on p.id = v.product_id
join public.catalog_brands b on b.id = p.brand_id and lower(b.name) = lower(s.brand);

do $validate$
begin
  if exists (select s.model from import_source s left join import_targets t
             on t.model = s.model group by s.model having count(t.variant_id) <> 1) then
    raise exception 'Import cancelled: each source model must match exactly one live variant';
  end if;
  if exists (select t.sku from import_targets t join public.catalog_variants v
             on lower(v.sku) = lower(t.sku) and v.id <> t.variant_id) then
    raise exception 'Import cancelled: a model SKU belongs to another variant';
  end if;
  if exists (select product_id from import_targets group by product_id
             having count(distinct name) > 1) then
    raise exception 'Import cancelled: conflicting names for a shared product';
  end if;
end;
$validate$;
"""
    photo_table = f"""create temporary table import_photos on commit drop as
select t.product_id, t.variant_id, t.name, t.brand, i.*
from jsonb_to_recordset({photos}::jsonb)
as i(model text, source text, path text, sha256 text, "originalBytes" integer,
     "webpBytes" integer, width integer, height integer, "isPrimary" boolean,
     "sortOrder" integer, "explicitCover" boolean, "storagePath" text)
join import_targets t on t.model = i.model;
"""
    backup = """select jsonb_build_object(
  'products', (select jsonb_agg(to_jsonb(p)) from public.catalog_products p
               where p.id in (select product_id from import_targets)),
  'variants', (select jsonb_agg(to_jsonb(v)) from public.catalog_variants v
               where v.id in (select variant_id from import_targets)),
  'images', (select jsonb_agg(to_jsonb(i)) from public.catalog_product_images i
             where i.product_id in (select product_id from import_targets))
) as backup;
"""
    (output / "inspect-live-catalog.sql").write_text(
        "begin;\n" + source_table + photo_table + backup + "rollback;\n")
    apply = """-- Apply only after uploading the files in upload/ to the catalog-products bucket.
-- This transaction preserves prices, inventory, publication, UUIDs and existing URLs.
begin;
""" + source_table + photo_table + """
do $storage$
begin
  if exists (select 1 from import_photos i where not exists (
    select 1 from storage.objects o where o.bucket_id = 'catalog-products'
      and 'catalog-products/' || o.name = i."storagePath"
      and o.metadata->>'mimetype' = 'image/webp')) then
    raise exception 'Import cancelled: one or more WebP objects have not been uploaded';
  end if;
end;
$storage$;

-- Keep a recoverable before-image in each affected product's existing metadata.
update public.catalog_products p
set metadata = p.metadata || jsonb_build_object('client_import_20261017_before',
  jsonb_build_object('name', p.name, 'description', p.description,
    'variants', (select jsonb_agg(to_jsonb(v)) from public.catalog_variants v where v.product_id = p.id),
    'images', (select jsonb_agg(to_jsonb(i)) from public.catalog_product_images i where i.product_id = p.id)))
where p.id in (select product_id from import_targets)
  and not (p.metadata ? 'client_import_20261017_before');

update public.catalog_products p
set name = t.name, description = t.specifications,
    metadata = p.metadata || jsonb_build_object('client_import_20261017',
      jsonb_build_object('workbook', 'Catalogo_Mosqueta_Productos_Actualizado.xlsx',
                        'row', t."sourceRow", 'model_as_sku', true))
from import_targets t where p.id = t.product_id;

update public.catalog_variants v
set sku = t.sku, manufacturer_model = t.model, source_model = t."sourceModel",
    color = t.color, specifications = t.specifications
from import_targets t where v.id = t.variant_id;

-- Retain old gallery records and files, after the final client photographs.
update public.catalog_product_images i
set is_primary = false, sort_order = greatest(i.sort_order, 100)
where i.product_id in (select distinct product_id from import_photos)
  and (i.variant_id is null or i.variant_id in (select variant_id from import_photos))
  and not exists (select 1 from import_photos f where f."storagePath" = i.storage_path);

insert into public.catalog_product_images
  (product_id, variant_id, storage_path, alt_text, image_status, is_primary, sort_order)
select product_id, variant_id, "storagePath",
       name || ' ' || brand || ' ' || model || case when "isPrimary" then ' - portada'
          else ' - imagen ' || ("sortOrder" + 1)::text end,
       'final', "isPrimary", "sortOrder"
from import_photos
on conflict (product_id, variant_id, storage_path) do update
set alt_text = excluded.alt_text, image_status = excluded.image_status,
    is_primary = excluded.is_primary, sort_order = excluded.sort_order;

select (select count(*) from import_targets) as updated_variants,
       (select count(*) from import_photos) as registered_images,
       (select count(distinct variant_id) from import_photos) as updated_galleries;
commit;
"""
    (output / "apply-catalog-update.sql").write_text(apply)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    workbook_paths = sorted(args.source.glob("*.xlsx"))
    if len(workbook_paths) != 1:
        raise ValueError("Expected exactly one source workbook")
    workbook_path = workbook_paths[0]
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Catálogo de Productos"]
    if list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))) != [
        "Artículo", "Marca", "Modelo", "Color", "Especificaciones"
    ]:
        raise ValueError("Unexpected workbook columns")

    source_rows = []
    for index, values in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
        if not any(value is not None for value in values):
            continue
        name, brand, model, color, specifications = values
        if not all(isinstance(value, str) and value.strip() for value in [name, brand, model, specifications]):
            raise ValueError(f"Missing required source data in row {index}")
        source_rows.append(dict(sourceRow=index, name=name.strip(), brand=brand.strip(),
                                sourceModel=model.strip(), color=color.strip() if color else None,
                                specifications=specifications.strip()))

    # Only spaced slashes separate alternative models. WK-5915/BD is one SKU.
    expanded = []
    for row in source_rows:
        for model in re.split(r"\s+/\s+", row["sourceModel"]):
            expanded.append({**row, "model": model, "sku": model})

    groups = defaultdict(list)
    for row in expanded:
        groups[(key(row["brand"]), key(row["model"]))].append(row)
    previous = json.loads((root / "src/data/products.generated.json").read_text())
    previous_groups = defaultdict(list)
    for row in previous:
        previous_groups[(key(row["brand"]), key(row["model"]))].append(row)

    ready, conflicts = [], []
    for group_key, rows in groups.items():
        candidates = previous_groups[group_key]
        if len(rows) != 1 or len(candidates) > 1:
            conflicts.append(dict(model=rows[0]["model"], sourceRows=rows,
                                  existing=[{k: r[k] for k in ["id", "name", "model", "category", "slug"]}
                                            for r in candidates],
                                  reason="duplicate-model"))
            continue
        ready.append({**rows[0], "existingSourceRef": candidates[0]["id"] if candidates else None,
                      "existingSlug": candidates[0]["slug"] if candidates else None,
                      "category": candidates[0]["category"] if candidates else None,
                      "match": "existing" if candidates else "new-model-review"})

    model_index = defaultdict(set)
    for row in expanded:
        model_index[key(row["model"])].add(row["model"])
    image_groups = defaultdict(list)
    for file in sorted(args.source.rglob("*")):
        if file.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        model_part = file.stem.split(".")[0].strip()
        candidates = model_index.get(key(model_part), set())
        if len(candidates) != 1:
            raise ValueError(f"Ambiguous or missing model for image {file.name}: {candidates}")
        model = next(iter(candidates))
        image_groups[model].append(file)

    args.output.mkdir(parents=True, exist_ok=True)
    images, previews = [], []
    for model, files in sorted(image_groups.items()):
        files.sort(key=lambda p: ("portada" not in p.stem.lower(), p.stem.lower(), str(p)))
        explicit_cover = [p for p in files if "portada" in p.stem.lower()]
        if len(explicit_cover) > 1:
            raise ValueError(f"Multiple covers for {model}")
        for position, file in enumerate(files):
            is_primary = position == 0
            suffix = "portada" if is_primary else f"{position:02d}"
            relative = Path("webp") / slug(model) / f"{slug(model)}-{suffix}.webp"
            dest = args.output / relative
            dest.parent.mkdir(parents=True, exist_ok=True)
            with Image.open(file) as source:
                oriented = ImageOps.exif_transpose(source)
                oriented.thumbnail((2000, 2000), Image.Resampling.LANCZOS)
                image = oriented.convert("RGB")
                image.save(dest, "WEBP", quality=82, method=6)
                width, height = image.size
            with Image.open(dest) as verify:
                verify.load()
                if verify.format != "WEBP" or verify.size != (width, height):
                    raise ValueError(f"Invalid WebP output: {dest}")
                thumb = verify.copy()
                thumb.thumbnail((230, 166), Image.Resampling.LANCZOS)
                previews.append((thumb, f"{model} / {suffix}"))
            digest = hashlib.sha256(dest.read_bytes()).hexdigest()
            upload_name = f"20261017-{slug(model)}-{suffix}-{digest[:12]}.webp"
            upload_dir = args.output / "upload"
            upload_dir.mkdir(exist_ok=True)
            shutil.copy2(dest, upload_dir / upload_name)
            images.append(dict(model=model, source=str(file.relative_to(args.source)),
                               path=relative.as_posix(), sha256=digest,
                               originalBytes=file.stat().st_size, webpBytes=dest.stat().st_size,
                               width=width, height=height, isPrimary=is_primary, sortOrder=position,
                               explicitCover=bool(explicit_cover) if is_primary else False,
                               storagePath=f"catalog-products/{upload_name}"))

    # A contact sheet is for reviewing the real photographs after compression.
    columns, cell_w, cell_h = 5, 250, 200
    preview = Image.new("RGB", (columns * cell_w, ((len(previews) + columns - 1) // columns) * cell_h), "#f1f5f9")
    draw = ImageDraw.Draw(preview)
    for index, (thumb, label) in enumerate(previews):
        x, y = (index % columns) * cell_w, (index // columns) * cell_h
        preview.paste(thumb, (x + (cell_w - thumb.width) // 2, y + (170 - thumb.height) // 2))
        draw.text((x + 10, y + 177), label, fill="#111827")
    preview.save(args.output / "contact-sheet.jpg", quality=90)

    summary = dict(sourceRows=len(source_rows), expandedRows=len(expanded),
                   uniqueModels=len(groups), readyModels=len(ready),
                   matchedModels=sum(r["match"] == "existing" for r in ready),
                   newModels=[r["model"] for r in ready if r["match"] != "existing"],
                   conflictingModels=[r["model"] for r in conflicts],
                   imageCount=len(images), modelsWithImages=len(image_groups),
                   originalBytes=sum(r["originalBytes"] for r in images),
                   webpBytes=sum(r["webpBytes"] for r in images))
    summary["reductionPercent"] = round(100 * (1 - summary["webpBytes"] / summary["originalBytes"]), 1)
    manifest = dict(sourceWorkbook=str(workbook_path), sourceWorkbookSha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                    summary=summary, ready=ready, conflicts=conflicts, images=images,
                    preserveFields=["price_mxn", "stock", "availability", "published", "slug"],
                    policy="Do not archive absent models. Do not resolve conflicting SKUs automatically.")
    (args.output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    prepare_sql(args.output, manifest)
    with zipfile.ZipFile(args.output / "imagenes-mosqueta-webp.zip", "w", zipfile.ZIP_STORED) as archive:
        for photo in images:
            archive.write(args.output / photo["path"], arcname=photo["path"])
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
